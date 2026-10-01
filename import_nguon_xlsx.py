# -*- coding: utf-8 -*-
"""Import nguồn cào từ Excel (hỗ trợ cả 2 vùng US và GLOBAL)
vào du_lieu_traffic/nguon_us.json và du_lieu_traffic/nguon_global.json.

Cấu trúc file Excel (như 'Nguồn cào (2).xlsx'):
- Phần US (Dòng 1 đến trước dòng phân tách Global):
    Col A: Key (NFL, WNBA, Chính trị, NCAA, vụ án, Chó mèo AI, Nascar,...)
    Col B: Page (link FB, username hoặc Tên Page)
    Col C: Subkey / Đội / Nhân vật
    Col D: Ghi chú
- Phần GLOBAL (Từ header 'LINK, PAGE, KEY' hoặc dòng bắt đầu Hoàng gia...):
    Col A: Link FB
    Col B: Tên Page
    Col C: Key / Chủ đề (Hoàng gia Thụy Điển, Đan Mạch, AFL,...)
    Col D: Ghi chú / Link phụ

Chạy: python import_nguon_xlsx.py [duong_dan_excel]
Mặc định tìm file mới nhất trong Downloads hoặc 'Nguồn cào link.xlsx'.
"""
import glob
import json
import os
import shutil
import sys
import time
import urllib.parse
from collections import Counter

sys.stdout.reconfigure(encoding="utf-8")
import openpyxl

DUONG_DAN = os.path.dirname(os.path.abspath(__file__))
THU_MUC_DATA = os.path.join(DUONG_DAN, "du_lieu_traffic")
FILE_US = os.path.join(THU_MUC_DATA, "nguon_us.json")
FILE_GLOBAL = os.path.join(THU_MUC_DATA, "nguon_global.json")


def clean_fb_url(u: str) -> str:
    if not u:
        return ""
    u = u.strip()
    if "profile.php" in u.lower():
        parsed = urllib.parse.urlparse(u)
        qs = urllib.parse.parse_qs(parsed.query)
        id_val = qs.get("id", [""])[0]
        if id_val:
            return f"https://www.facebook.com/profile.php?id={id_val}"
        return u.rstrip("/")
    return u.split("?")[0].rstrip("/")


def tim_file_excel_mac_dinh() -> str:
    user_home = os.path.expanduser("~")
    downloads = os.path.join(user_home, "Downloads")
    candidates = glob.glob(os.path.join(downloads, "*nguồn*cào*.xlsx")) + \
                 glob.glob(os.path.join(downloads, "*nguon*cao*.xlsx"))
    if candidates:
        candidates.sort(key=os.path.getmtime, reverse=True)
        return candidates[0]
    local_f = os.path.join(DUONG_DAN, "Nguồn cào link.xlsx")
    if os.path.isfile(local_f):
        return local_f
    return ""


def import_nguon_tu_excel(fp_xlsx: str) -> dict:
    if not os.path.isfile(fp_xlsx):
        raise FileNotFoundError(f"Không tìm thấy file: {fp_xlsx}")

    print(f"📂 Đang đọc dữ liệu từ: {fp_xlsx}")
    wb = openpyxl.load_workbook(fp_xlsx, data_only=True)
    ws = wb[wb.sheetnames[0]]

    # Đọc dữ liệu cũ để tra cứu tên hiển thị -> link FB
    old_us = []
    if os.path.isfile(FILE_US):
        try:
            with open(FILE_US, "r", encoding="utf-8") as f:
                old_us = json.load(f)
        except Exception:
            pass

    us_name_to_entry = {}
    for x in old_us:
        tp = (x.get("Tên page") or "").strip().lower()
        if tp:
            us_name_to_entry[tp] = x

    # 1. Nhận diện dòng phân cách giữa US và Global
    row_split = None
    for r in range(1, ws.max_row + 1):
        a = str(ws.cell(r, 1).value or "").strip().upper()
        b = str(ws.cell(r, 2).value or "").strip().upper()
        c = str(ws.cell(r, 3).value or "").strip().upper()
        if (a == "LINK" and b == "PAGE") or "HOÀNG GIA" in c or "HOÀNG GIA" in a:
            row_split = r
            break

    if row_split is None:
        row_split = ws.max_row + 1

    print(f"📍 Dòng phân tách Global: Row {row_split}")

    # 2. Xử lý US (từ dòng 2 đến row_split - 1)
    us_list = []
    us_seen = set()
    curr_key = ""

    for r in range(2, row_split):
        k = ws.cell(r, 1).value
        p = str(ws.cell(r, 2).value or "").strip()
        c = str(ws.cell(r, 3).value or "").strip()
        note = str(ws.cell(r, 4).value or "").strip()

        if k is not None and str(k).strip():
            curr_key = str(k).strip()
        if not p:
            continue

        entry = None
        if p.lower().startswith("http") or "facebook.com" in p.lower() or p.lower().startswith("profile.php"):
            url = clean_fb_url(p)
            name = ""
        elif " " not in p:
            url = clean_fb_url(f"https://www.facebook.com/{p}")
            name = p
        else:
            entry = us_name_to_entry.get(p.lower())
            url = clean_fb_url(entry.get("Facebook nguồn", "")) if entry else ""
            name = p

        if entry and not name:
            name = entry.get("Tên page", "")

        # Phân loại Key và subkey môn thể thao
        if c:
            if curr_key.upper() == "NFL":
                c_low = c.lower()
                if any(x in c_low for x in ["tigers", "rebels", "gamecocks", "yellow jackets"]) and "clemson" not in c_low and "lsu" not in c_low:
                    key_final = f"NCAA - {c}"
                elif any(x in c_low for x in ["avalanche", "oilers"]):
                    key_final = f"NHL - {c}"
                elif "astros" in c_low:
                    key_final = f"MLB - {c}"
                else:
                    key_final = f"NFL - {c}"
            elif curr_key.upper() == "NCAA":
                key_final = f"NCAA - {c}"
            elif curr_key.upper() == "MLB":
                key_final = f"MLB - {c}"
            else:
                key_final = curr_key
        else:
            key_final = curr_key

        if key_final.lower() == "vụ án":
            key_final = "Vụ án"
        elif key_final.lower() == "phim cao bồi":
            key_final = "Phim cao bồi"
        elif key_final.lower() == "nghệ sĩ đồng quê":
            key_final = "Nghệ sĩ đồng quê"

        if url and url.lower() not in us_seen:
            us_seen.add(url.lower())
            item = {
                "KEY": key_final,
                "Facebook nguồn": url,
                "Vung": "us"
            }
            if name:
                item["Tên page"] = name
            nv = c or (entry.get("Nhân vật gợi ý", "") if entry else "")
            if nv:
                item["Nhân vật gợi ý"] = nv
            gh = note or (entry.get("Ghi chú", "") if entry else "")
            if gh:
                item["Ghi chú"] = gh

            us_list.append(item)

    # 3. Xử lý GLOBAL (từ row_split đến hết)
    global_list = []
    global_seen = set()
    curr_key_gb = ""

    start_gb = row_split + 1 if str(ws.cell(row_split, 1).value or "").strip().upper() == "LINK" else row_split

    for r in range(start_gb, ws.max_row + 1):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        c = ws.cell(r, 3).value
        d = ws.cell(r, 4).value

        if c is not None and str(c).strip():
            curr_key_gb = str(c).strip()

        url = ""
        name = str(b or "").strip()
        note = ""

        if a and str(a).strip().startswith("http"):
            url = clean_fb_url(str(a).strip())
        elif d and str(d).strip().startswith("http"):
            url = clean_fb_url(str(d).strip())
        elif a and not str(a).strip().startswith("http") and " " not in str(a).strip():
            url = clean_fb_url(f"https://www.facebook.com/{str(a).strip()}")

        if not url:
            continue

        key_clean = curr_key_gb
        if "\n" in key_clean:
            parts = key_clean.split("\n")
            key_clean = parts[0].strip()
            note = "\n".join(parts[1:]).strip()

        if key_clean.lower() == "tenis":
            key_clean = "Tennis"
        elif "bóng đá" in key_clean.lower():
            key_clean = "Bóng đá"
        elif key_clean.lower() == "hoàng gua + vụ án đan mạch":
            key_clean = "Hoàng gia + Vụ án Đan Mạch"
        elif key_clean.lower() == "hoàng gia vói người nổi tiếng":
            key_clean = "Hoàng gia với người nổi tiếng"
        elif key_clean.lower() == "hoàng gia và đua xe dạp":
            key_clean = "Hoàng gia và đua xe đạp"
        elif key_clean.lower() == "hoàng gia nauy":
            key_clean = "Hoàng gia Na Uy"

        if url and url.lower() not in global_seen:
            global_seen.add(url.lower())
            item = {
                "KEY": key_clean,
                "Facebook nguồn": url,
                "Vung": "global"
            }
            if name:
                item["Tên page"] = name
            if note:
                item["Ghi chú"] = note
            global_list.append(item)

    # 4. Tạo bản sao lưu backup
    ts = int(time.time())
    if os.path.isfile(FILE_US):
        shutil.copy2(FILE_US, f"{FILE_US}.bak_{ts}")
    if os.path.isfile(FILE_GLOBAL):
        shutil.copy2(FILE_GLOBAL, f"{FILE_GLOBAL}.bak_{ts}")

    # 5. Ghi file JSON cho US và GLOBAL
    os.makedirs(THU_MUC_DATA, exist_ok=True)
    with open(FILE_US, "w", encoding="utf-8") as f:
        json.dump(us_list, f, ensure_ascii=False, indent=2)

    with open(FILE_GLOBAL, "w", encoding="utf-8") as f:
        json.dump(global_list, f, ensure_ascii=False, indent=2)

    print(f"\n🎉 CẬP NHẬT THÀNH CÔNG:")
    print(f"   🇺🇸 US ({FILE_US}): {len(us_list)} nguồn cào")
    c_us = Counter(x["KEY"] for x in us_list)
    for k, v in sorted(c_us.items()):
        print(f"      - {k:35s}: {v} nguồn")

    print(f"\n   🌐 GLOBAL ({FILE_GLOBAL}): {len(global_list)} nguồn cào")
    c_gb = Counter(x["KEY"] for x in global_list)
    for k, v in sorted(c_gb.items()):
        print(f"      - {k:35s}: {v} nguồn")

    return {
        "ok": True,
        "so_us": len(us_list),
        "so_global": len(global_list),
        "file_us": FILE_US,
        "file_global": FILE_GLOBAL,
    }


def main():
    fp_xlsx = sys.argv[1] if len(sys.argv) > 1 else tim_file_excel_mac_dinh()
    if not fp_xlsx:
        print("❌ Không tìm thấy file Excel nguồn cào trong Downloads hoặc thư mục dự án.")
        return 1
    import_nguon_tu_excel(fp_xlsx)
    return 0


if __name__ == "__main__":
    sys.exit(main())
